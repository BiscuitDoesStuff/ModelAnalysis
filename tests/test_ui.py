"""Phase 5 UX/accessibility: shared UI layer, page structure, zero vs unknown, cross-format rows."""
import contextlib
import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "tests"), str(ROOT / "tests" / "fixtures")]
import make_golden  # noqa: E402
import pipeline  # noqa: E402
from reports import build_report, ui  # noqa: E402
from reports.build_site import model_filename  # noqa: E402
from test_site import Document, fixture, generate  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures"


def cell_text(node):
    """Visible text, with a price cell's source hint (after <br>) separated by one space."""
    hints = [c.text for c in node.children if c.has_class("hint")]
    text = node.text[:len(node.text) - sum(map(len, hints))] if hints else node.text
    return " ".join(" ".join([text, *hints]).split())


class UnitTests(unittest.TestCase):
    def test_contrast_formula(self):
        self.assertAlmostEqual(ui.contrast("#000000", "#ffffff"), 21.0, places=2)
        self.assertAlmostEqual(ui.contrast("#777777", "#ffffff"), 4.48, places=2)
        self.assertEqual(ui.contrast("#123456", "#123456"), 1.0)

    def test_token_pairs_meet_wcag_aa_in_both_themes(self):
        self.assertEqual(set(ui.TOKENS), {"dark", "light"})
        self.assertEqual(set(ui.TOKENS["dark"]), set(ui.TOKENS["light"]))
        for theme, tokens in ui.TOKENS.items():
            for fg, bg in ui.CONTRAST_PAIRS:
                ratio = ui.contrast(tokens[fg], tokens[bg])
                self.assertGreaterEqual(ratio, ui.AA_TEXT, f"{theme}: {fg} on {bg} is {ratio:.2f}")

    def test_css_has_light_theme_focus_and_narrow_layout(self):
        css = ui.css()
        for needle in ("prefers-color-scheme: light", ":focus-visible", "max-width:480px", ".skip:focus", ".unk"):
            self.assertIn(needle, css)
        self.assertNotIn("opacity", css)  # dimming text breaks the contrast guarantee

    def test_copy_button_is_named_and_empty_for_display_only(self):
        doc = Document(ui.copy_button("openrouter/a/b<c>"))
        button = doc.nodes[0]
        self.assertEqual((button.tag, button.attrs["type"]), ("button", "button"))
        self.assertEqual(button.attrs["data-copy"], "openrouter/a/b<c>")
        self.assertEqual(button.attrs["aria-label"], "Copy route openrouter/a/b<c>")
        self.assertEqual(ui.copy_button(""), "")

    def test_tabs_wiring(self):
        doc = Document(ui.tabs([("p1", "One", "a"), ("p2", "Two", "b")], "Pages"))
        tabs = [n for n in doc.nodes if n.attrs.get("role") == "tab"]
        self.assertEqual([t.attrs["aria-selected"] for t in tabs], ["true", "false"])
        self.assertEqual([t.attrs["tabindex"] for t in tabs], ["0", "-1"])
        self.assertNotIn("hidden", doc.by_id("p1").attrs)
        self.assertIn("hidden", doc.by_id("p2").attrs)
        self.assertEqual(ui.check_html(ui.tabs([("p1", "One", "a")], "Pages")), [])

    def test_check_html_catches_each_problem(self):
        cases = {
            "<button type='button'></button>": "button without accessible name",
            "<input type='search'>": "<input> without label",
            "<select id='s'></select>": "<select> without label",
            "<table><tr><td>x</td></tr></table>": "table without caption",
            "<table><caption>c</caption><tr><th>x</th></tr></table>": "th without scope",
            "<button role='tab' id='t' aria-controls='nope'>T</button>": "controls no tabpanel",
            "<div role='tabpanel' id='p' aria-labelledby='x'></div>": "has no tab",
            "<code class='copy'>openrouter/x</code>": "not a button",
            "<button class='copy' data-copy='display/only'>display/only</button>": "non-callable route",
        }
        for page, problem in cases.items():
            found = ui.check_html(page, allowed_routes={"openrouter/x"})
            self.assertTrue(any(problem in p for p in found), f"{page!r}: {found}")
        clean = ("<label for='s'>S</label><select id='s'></select><label>F <input></label>"
                 "<input aria-label='x'><button aria-label='Close'></button>"
                 "<table><caption>c</caption><tr><th scope='col'>x</th></tr></table>"
                 "<button class='copy' data-copy='openrouter/x'>openrouter/x</button>")
        self.assertEqual(ui.check_html(clean, allowed_routes={"openrouter/x"}), [])


class ZeroVersusUnknownTests(unittest.TestCase):
    def test_zero_price_and_unknown_price_render_differently(self):
        a = fixture(count=4)
        a["models"][2]["cost_blended"] = None
        a["models"][3]["score"] = None
        with tempfile.TemporaryDirectory() as tmp:
            out, site = generate(Path(tmp), a)
            pages = {"dashboard": out / f"{a['stamp']}_report.html",
                     "site": site / "models.html"}
            for name, path in pages.items():
                doc = Document(path.read_text(encoding="utf-8"))
                rows = {}
                for tr in (n for n in doc.nodes if n.tag == "tr"):
                    code = next((n for n in doc.nodes if n.tag == "code" and n.parent and self.within(n, tr)), None)
                    if code:
                        rows.setdefault(code.text, tr)
                zero = next(c for c in rows["vendor/model-0"].children if c.has_class("cc"))
                unknown = next(c for c in rows["vendor/model-2"].children if c.has_class("cc"))
                unscored = next(c for c in rows["vendor/model-3"].children if c.has_class("sc"))
                self.assertTrue(cell_text(zero).startswith("$0/1M"), name)
                self.assertFalse(any(n.has_class("unk") for n in zero.children), name)
                self.assertEqual([n.text for n in unknown.children if n.has_class("unk")], ["unknown"], name)
                self.assertEqual([n.text for n in unscored.children if n.has_class("unk")], ["unscored"], name)

    @staticmethod
    def within(node, root):
        while node.parent:
            node = node.parent
            if node is root:
                return True
        return False


class GoldenBundleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        root = Path(cls.temp.name)
        with contextlib.redirect_stdout(io.StringIO()):
            cls.bundle = pipeline.run(state_dir=root / "runs", db_path=root / "h.sqlite",
                                      snapshot_path=FIXTURES / "golden_snapshot.json",
                                      websites_path=FIXTURES / "golden_websites.json",
                                      registry_path=FIXTURES / "golden_research.json", as_of=make_golden.AS_OF)
        cls.run_id = json.loads((cls.bundle / "manifest.json").read_text(encoding="utf-8"))["run_id"]
        cls.reports = cls.bundle / "reports"
        cls.site = cls.reports / f"{cls.run_id}_site"
        cls.dashboard = cls.reports / f"{cls.run_id}_report.html"
        cls.a = json.loads((cls.bundle / "analysis" / f"{cls.run_id}_analysis.json").read_text(encoding="utf-8"))
        cls.models = [m for m in cls.a["models"] if not m.get("router")]

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def pages(self):
        return [self.dashboard, *sorted(self.site.rglob("*.html"))]

    def test_every_page_is_structurally_clean(self):
        for path in self.pages():
            text = path.read_text(encoding="utf-8")
            self.assertEqual(ui.check_html(text), [], path.name)
            doc = Document(text)
            self.assertEqual(doc.errors, [], path.name)
            self.assertEqual(doc.nodes[0].tag if doc.nodes else None, "html")
            self.assertEqual(doc.by_id("main").tag, "main")
            skip = next(n for n in doc.nodes if n.tag == "a")
            self.assertEqual((skip.attrs.get("href"), skip.attrs.get("class")), ("#main", "skip"))
            self.assertEqual(doc.by_id("copy-status").attrs.get("aria-live"), "polite")
            for tablist in (n for n in doc.nodes if n.attrs.get("role") == "tablist"):
                self.assertTrue(tablist.attrs.get("aria-label"))
                selected = [t for t in tablist.children if t.attrs.get("aria-selected") == "true"]
                self.assertEqual(len(selected), 1, path.name)

    def test_dashboard_tabs_and_graph_are_named(self):
        doc = Document(self.dashboard.read_text(encoding="utf-8"))
        tabs = [n for n in doc.nodes if n.attrs.get("role") == "tab"]
        self.assertEqual([t.attrs["aria-controls"] for t in tabs],
                         ["t-start", "t-value", "t-stack", "t-compare", "t-free", "t-graph", "t-explore"])
        plot = doc.by_id("g-plot")
        self.assertEqual(plot.attrs["role"], "img")
        self.assertEqual(plot.attrs["aria-labelledby"], "g-plot-title g-plot-desc")
        self.assertEqual(doc.by_id("g-plot-title").tag, "title")
        self.assertEqual(doc.by_id("g-plot-desc").tag, "desc")
        js = (ROOT / "reports" / "graph.js").read_text(encoding="utf-8")
        # renderPlot clears the SVG, so it must re-add both, and the fallback table is captioned.
        self.assertIn("svg('title', {id: 'g-plot-title'}", js)
        self.assertIn("svg('desc', {id: 'g-plot-desc'}", js)
        self.assertIn("element('caption'", js)

    def test_display_only_routes_never_get_a_copy_button(self):
        display_only = [m for m in self.models if not build_report.copy_id(m)]
        self.assertTrue(display_only, "golden fixture should have display-only rows")
        for m in display_only:
            doc = Document((self.site / "models" / model_filename(m["slug"])).read_text(encoding="utf-8"))
            self.assertEqual([n for n in doc.nodes if n.has_class("copy") and not self.route_row(n)], [], m["id"])
            if m.get("nearest_callable"):
                self.assertNotIn(f"data-copy='{m['nearest_callable']}'", self.dashboard.read_text(encoding="utf-8"))

    @staticmethod
    def route_row(node):
        """Per-provider route rows (Routes table) copy their own catalog IDs."""
        while node.parent:
            node = node.parent
            if node.tag == "table":
                return any(c.tag == "caption" and c.text == "Provider routes" for c in node.children)
        return False

    def dashboard_row(self, model_id):
        doc = Document(self.dashboard.read_text(encoding="utf-8"))
        tbody = next(n for n in doc.by_id("xtab").children if n.tag == "tbody")
        row = next(tr for tr in tbody.children
                   if next(c for c in tr.children[0].children if c.tag == "code").text == model_id)
        cells = [c for c in row.children if c.tag == "td"]
        headers = build_report.MODEL_HEADERS
        route = cells[headers.index("Route ID / selector (copy)")]
        button = next((n for n in route.children if n.tag == "button"), None)
        return {"score": cell_text(cells[headers.index("Score")]),
                "price": cell_text(cells[headers.index("Price")]),
                "free": cell_text(cells[headers.index("Free")]),
                "route": button.attrs["data-copy"] if button else ""}

    def site_row(self, m):
        doc = Document((self.site / "models" / model_filename(m["slug"])).read_text(encoding="utf-8"))
        rows = {}
        for tr in (n for n in doc.nodes if n.tag == "tr"):
            th = next((c for c in tr.children if c.tag == "th" and c.attrs.get("scope") == "row"), None)
            if th:
                rows[th.text] = [c for c in tr.children if c.tag == "td"]
        route_p = next(n for n in doc.nodes if n.tag == "p" and n.text.startswith("Route:"))
        button = next((n for n in route_p.children if n.tag == "button"), None)
        summary = next(n for n in doc.nodes if n.tag == "p")
        return {"score": cell_text(rows["AA Intelligence"][0]),
                "price": cell_text(rows["Price (USD/1M, 3:1 blend)"][0]),
                "free": summary.text.rsplit(" · ", 1)[-1],
                "route": button.attrs["data-copy"] if button else ""}

    def xlsx_row(self, model_id):
        from openpyxl import load_workbook
        wb = load_workbook(self.reports / f"{self.run_id}_models.xlsx", read_only=True)
        try:
            rows = wb["All_Intel"].iter_rows(values_only=True)
            header = list(next(rows))
            row = dict(zip(header, next(r for r in rows if r[header.index("id")] == model_id)))
        finally:
            wb.close()
        suffix = " (estimate)" if row["score_source"].startswith("inherited estimate") else ""
        price = "unknown" if row["cost_per_1M"] is None else f"${row['cost_per_1M']}/1M"
        return {"score": "unscored" if row["score"] is None else f"{row['score']}{suffix}",
                "price": price + (f" {row['cost_source'] or 'none'}" if row["cost_per_1M"] is not None else ""),
                "free": build_report.FREE_LABEL[row["free_status"]],
                "route": row["opencode_id"] or ""}

    def test_common_rows_match_across_dashboard_site_and_xlsx(self):
        picks = {
            "scored, costed, callable": next(m for m in self.models if m.get("score") is not None
                                              and m.get("cost_blended") is not None and build_report.copy_id(m)),
            "uncosted": next(m for m in self.models if m.get("cost_blended") is None),
            "display-only": next(m for m in self.models if not build_report.copy_id(m)),
            "provisional free": next(m for m in self.models if build_report.is_provisional(m)),
        }
        # XLSX cells come back as int or float; compare the numbers, not their spelling.
        def numeric(row):
            return {k: re.sub(r"\d+(?:\.\d+)?", lambda x: repr(float(x.group())), v) if k in ("score", "price") else v
                    for k, v in row.items()}

        for label, m in picks.items():
            view = build_report.row_view(m)
            expected = {"score": view["score"],
                        "price": view["price"] + (f" {view['price_source']}" if view["price_known"] else ""),
                        "free": view["free"], "route": view["route"]}
            for fmt, got in (("dashboard", self.dashboard_row(m["id"])), ("site", self.site_row(m)),
                             ("xlsx", self.xlsx_row(m["id"]))):
                self.assertEqual(numeric(got), numeric(expected), f"{label} {m['id']} in {fmt}")


if __name__ == "__main__":
    unittest.main()
