"""Smoke check for the 9-section report contract + Phase-1 provisional-free. Run: python tests/smoke.py"""
import json, glob, os, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
fails = []


def check(cond, msg):
    print(("ok  " if cond else "FAIL") + f" {msg}")
    if not cond:
        fails.append(msg)


a_files = sorted(glob.glob(os.path.join(ROOT, "analysis", "*_analysis.json")), key=os.path.getmtime)
r_files = sorted(glob.glob(os.path.join(ROOT, "reports", "*_models.json")), key=os.path.getmtime)
h_files = sorted(glob.glob(os.path.join(ROOT, "reports", "*_report.html")), key=os.path.getmtime)
m_files = sorted(glob.glob(os.path.join(ROOT, "reports", "*_summary.md")), key=os.path.getmtime)
check(a_files and r_files, "analysis + report files exist")
check(bool(h_files), "html dashboard exists")
if not (a_files and r_files):
    sys.exit(1)

a = json.load(open(a_files[-1], encoding="utf-8"))
r = json.load(open(r_files[-1], encoding="utf-8"))
check(a.get("stamp") == r.get("stamp"), f"stamps match ({a.get('stamp')})")
check(r.get("thresholds") == {"max": 50, "high": 40, "medium": 30}, "thresholds 50/40/30")
check(len(a.get("models", [])) > 0, f"canonical models non-empty ({len(a.get('models', []))})")

for key in ["all_intel", "all_cost", "all_ratio_paid", "all_ratio_free_by_score",
            "ocf_intel", "ocf_cost", "ocf_ratio_paid",
            "ocf_stack", "ocf_practical", "ocf_outliers"]:
    check(key in r, f"report key {key}")

# Phase-1 provisional-free keys (third ratio block).
for key in ["all_ratio_verified_free_by_score", "all_ratio_provisional_free_by_score",
            "ocf_ratio_verified_free_by_score", "ocf_ratio_provisional_free_by_score",
            "free_status_counts"]:
    check(key in r or key in a, f"phase1 key {key}")

tiers = set(a.get("thresholds", {})) | {"max", "high", "medium"}
check(set(r["ocf_stack"]) == {"max", "high", "medium"}, "stack has 3 tiers")
for t in ("max", "high", "medium"):
    rows = r["ocf_stack"][t]["rows"]
    check(all(m.get("tier") == t for m in rows), f"stack[{t}] tiers consistent ({len(rows)})")
    check(set(r["ocf_stack"][t]["gaps"]) <= {"O", "C", "F"}, f"stack[{t}] gaps valid")
check(len(r["ocf_practical"]) == 12, f"practical has 12 tier x variant rows ({len(r['ocf_practical'])})")
check({(p["tier"], p["variant"]) for p in r["ocf_practical"]} ==
      {(t, v) for t in ("max", "high", "medium") for v in ("OCF", "OF", "CF", "F")},
      "practical covers all tier x variant combos")
check(set(r["ocf_outliers"]) == {"bargains", "overpriced", "free_gems"}, "outlier classes")

# Phase-1: free_status values + evidence.
allowed = {"verified", "provisional-l1", "provisional-l0", "none"}
models = a.get("models", [])
check(all(m.get("free_status") in allowed for m in models), "free_status values valid")
check(all(isinstance(m.get("free_evidence"), list) for m in models), "free_evidence lists present")
check(all((m.get("free") is True) == (m.get("free_status") == "verified") for m in models),
      "free bool matches verified status")
check(all(("F" in m.get("groups", [])) == (m.get("free_status") == "verified") for m in models
          if not m.get("router")),
      "F group stays verified-only")
# L1 must have an OR listing; L0 must have no OR listing (may carry groq/cerebras native).
l1 = [m for m in models if m.get("free_status") == "provisional-l1"]
l0 = [m for m in models if m.get("free_status") == "provisional-l0"]
check(all(m.get("or_id") for m in l1), f"L1 all have or_id ({len(l1)})")
check(all(not m.get("or_id") for m in l0),
      f"L0 all lack or_id ({len(l0)})")
fsc = a.get("free_status_counts", {})
check(sum(fsc.get(k, 0) for k in ("verified", "provisional-l1", "provisional-l0", "none")) == len(models),
      f"free_status_counts sum to models ({fsc})")

# Phase-1: OCF includes provisional; stack/practical/outliers need callable IDs.
ocf_ids = {m["id"] for m in r.get("ocf_intel", [])}
check(all(m["id"] in ocf_ids for m in l1 if not m.get("router")),
      "L1 callable provisional in OCF intel")
check(all(m["id"] in ocf_ids for m in l0 if not m.get("router")),
      "L0 AA-only provisional in OCF intel (confined to Intel/Cost/Ratio)")
for t in ("max", "high", "medium"):
    for m in r["ocf_stack"][t]["rows"]:
        check(bool(m.get("or_id")) or bool(set(m.get("providers", [])) & {"openai", "anthropic", "openrouter", "groq", "cerebras"}),
              f"stack[{t}] callable only")
        if not (bool(m.get("or_id")) or bool(set(m.get("providers", [])) & {"openai", "anthropic", "openrouter", "groq", "cerebras"})):
            break
# Ratio three blocks: paid / verified-free / provisional-free, no overlap.
paid_ids = {m["id"] for m in r.get("all_ratio_paid", [])}
ver_ids = {m["id"] for m in r.get("all_ratio_verified_free_by_score", r.get("all_ratio_free_by_score", []))}
prov_ids = {m["id"] for m in r.get("all_ratio_provisional_free_by_score", [])}
check(not (paid_ids & ver_ids), "paid vs verified-free disjoint")
check(not (paid_ids & prov_ids), "paid vs provisional-free disjoint")
check(not (ver_ids & prov_ids), "verified vs provisional disjoint")

# Phase-2: Groq + Cerebras catalogs (keyed-and-skipped like OAI/ANT).
check(isinstance(a.get("total_groq"), int), f"total_groq present ({a.get('total_groq')})")
check(isinstance(a.get("total_cerebras"), int), f"total_cerebras present ({a.get('total_cerebras')})")
check(isinstance(a.get("groq_ids"), list), "groq_ids list present")
check(isinstance(a.get("cerebras_ids"), list), "cerebras_ids list present")
check(all(set(m.get("providers", [])) <= {"openrouter", "openai", "anthropic", "groq", "cerebras"}
          for m in models),
      "providers tags valid")
check(all(m.get("slug") for m in models), "canonical slugs present")
w_files = sorted(glob.glob(os.path.join(ROOT, "raw", "*_models.json")), key=os.path.getmtime)
if w_files:
    snap = json.load(open(w_files[-1], encoding="utf-8"))
    for src in ("groq", "cerebras"):
        v = snap.get(src, "MISSING")
        check(isinstance(v, list) or (isinstance(v, dict) and ("skipped" in v or "error" in v)),
              f"snapshot {src} list-or-skipped")

# Phase-3: free-status churn tracking.
ch = a.get("free_churn") or {}
check(isinstance(a.get("free_churn"), dict), "analysis free_churn present")
for key in ["flipped_to_paid", "flipped_to_paid_total", "flipped_to_free", "flipped_to_free_total",
            "disappeared", "disappeared_total", "new_slugs", "new_total",
            "or_flipped_to_paid", "or_flipped_to_paid_total",
            "or_flipped_to_free", "or_flipped_to_free_total"]:
    check(key in ch, f"churn key {key}")
check(all(isinstance(ch.get(k), int) for k in
          ["flipped_to_paid_total", "flipped_to_free_total", "disappeared_total", "new_total"]),
      "churn totals are ints")
check(all(len(ch.get(k, [])) <= 50 for k in
          ["flipped_to_paid", "flipped_to_free", "disappeared", "new_slugs"]),
      "churn lists capped at 50")
check("free_churn" in r, "report carries free_churn")
if m_files:
    md = open(m_files[-1], encoding="utf-8").read()
    check("[F?]" in md, "md has [F?] marker")
if h_files:
    htm = open(h_files[-1], encoding="utf-8").read()
    check("F?" in htm, "html has F? marker")

try:
    from openpyxl import load_workbook
    wb = load_workbook(r_files[-1].replace("_models.json", "_models.xlsx"))
    want = {"summary", "All_Intel", "All_Cost", "All_Ratio", "OCF_Intel",
            "OCF_Cost", "OCF_Ratio", "OCF_Outliers", "OCF_Stack", "OCF_Practical"}
    check(set(wb.sheetnames) == want, f"xlsx has 10 tabs ({len(wb.sheetnames)})")
    hdr = [c.value for c in wb["All_Ratio"][1]]
    check("free_status" in hdr, "xlsx All_Ratio has free_status column")
except ImportError:
    print("skip xlsx check (openpyxl missing)")

print(f"{len(fails)} failures")
sys.exit(1 if fails else 0)
