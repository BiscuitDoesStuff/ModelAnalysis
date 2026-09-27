"""Offline report/site integration checks. All generated files live in temp dirs."""
import contextlib
import io
import json
from html.parser import HTMLParser
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import unquote, urlsplit

from reports import build_report, build_site


class Node:
    def __init__(self, tag, attrs=(), parent=None):
        self.tag = tag
        self.attrs = dict(attrs)
        self.parent = parent
        self.children = []
        self.text = ""

    def has_class(self, name):
        return name in self.attrs.get("class", "").split()


class Document(HTMLParser):
    VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]
        self.nodes = []
        self.errors = []
        self.feed(text)
        self.close()
        if len(self.stack) != 1:
            self.errors.append("Unclosed: " + ", ".join(n.tag for n in self.stack[1:]))

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.stack[-1])
        self.stack[-1].children.append(node)
        self.nodes.append(node)
        if tag not in self.VOID:
            self.stack.append(node)

    def handle_endtag(self, tag):
        if len(self.stack) == 1 or self.stack[-1].tag != tag:
            self.errors.append(f"Unexpected </{tag}> inside <{self.stack[-1].tag}>")
        else:
            self.stack.pop()

    def handle_data(self, data):
        for node in self.stack:
            node.text += data

    def by_id(self, value):
        return next(n for n in self.nodes if n.attrs.get("id") == value)


def fixture(count=325, metadata=True):
    models = []
    for i in range(count):
        # Put the uncallable free pick first to exercise dashboard cards too.
        model = {"slug": f"model-{i}", "id": f"vendor/model-{i}", "name": f"Model {i}",
                 "score": 60 - i / 100, "cost_blended": 0 if i == 0 else 1,
                 "ratio": None if i == 0 else 40, "groups": ["F"] if i == 0 else ["O"],
                 "tier": "max", "providers": ["aa"] if i % 2 == 0 else ["openrouter"],
                 "free_status": "provisional-l0" if i == 0 else "none",
                 "or_id": "" if i % 2 == 0 else f"vendor/model-{i}",
                 "score_source": {"kind": "measured"}}
        models.append(model)
    if count:
        models[-1]["slug"] = "model/with <special>#%"
        models[0]["fallback_id"] = "aa-display-only"
        models[0]["fallback_provider"] = "aa"
    views = {
        "benchlm_leaderboard": [
            {"slug": m["slug"], "id": m["id"], "model": m["name"], "overall": i,
             "evidence": "estimated" if i < 110 else "supported", "categories": {"coding": i}}
            for i, m in enumerate(models)],
        "llmstats_leaderboard": [{"slug": m["slug"], "id": m["id"], "score": 40, "url": "https://example.test/model"}
                                 for m in models[-2:]],
        "vals_leaderboard": [{"slug": m["slug"], "id": m["id"], "accuracy": 50} for m in models[-2:]],
        "provisional_triage": [{"id": m["id"], "route": "display-only", "qualifier": False} for m in models[:2]],
    }
    a = {"stamp": "2026-09-26_fixture", "day": "2026-09-26", "models": models,
         "views": views}
    if count:
        a["models"] = models + [{**models[1], "slug": "router", "id": "router", "router": True}]
    if metadata:
        a.update({"run_id": "run-fixture", "schema_version": 3, "started_at": "2026-09-26T12:00:00Z",
                  "source_health": {"openrouter": {"status": "complete", "count": count, "complete": True,
                                                   "fetched_at": "2026-09-26T12:01:00Z"},
                                    "aa": {"status": "failed", "error": "<offline> & unavailable"}},
                  "website_health": {"vals": {"status": "partial", "complete": False, "count": 2,
                                              "scope": "selected-pages", "reason": "page unavailable",
                                              "components": {"pages": {"attempted": 3, "failed": 1}}}},
                  "churn": {"events": [{"kind": "disappeared", "source": "openrouter",
                                         "route_id": "vendor/old", "model": "Old model",
                                         "alternatives": ["vendor/new"], "loss_reason": "catalog removal", "coverage": "partial"}],
                            "baselines": {"openrouter": {"run_id": "prior-run", "day": "2026-09-25"}},
                            "daily": {"2026-09-26": {"disappeared": 1}}}})
    return a


def generate(root, a):
    source = root / "input.json"
    out = root / "reports"
    out.mkdir()
    source.write_text(json.dumps(a), encoding="utf-8")
    # Existing artifacts must survive explicit runs, regardless of timestamps.
    (out / "old_models.json").write_text("keep", encoding="utf-8")
    (out / "old_site").mkdir()
    (out / "old_site" / "keep.txt").write_text("keep", encoding="utf-8")
    with contextlib.redirect_stdout(io.StringIO()):
        build_report.main(source, out)
        build_site.main(source, out)
    return out, out / (a["stamp"] + "_site")


class SiteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.a = fixture()
        cls.out, cls.site = generate(Path(cls.temp.name), cls.a)

    def test_all_models_and_local_links_close(self):
        expected = [m for m in self.a["models"] if not m.get("router")]
        pages = list((self.site / "models").glob("*.html"))
        self.assertEqual(len(pages), len(expected))
        directory = Document((self.site / "models.html").read_text(encoding="utf-8"))
        links = [n for n in directory.nodes if n.tag == "a" and n.attrs.get("href", "").startswith("models/")]
        self.assertEqual({n.text for n in links}, {m["id"] for m in expected})
        # The directory filter sits in the same .tblock as the full table, so it searches every model.
        search = next(n for n in directory.nodes if n.tag == "input" and "data-filter" in n.attrs)
        block = search
        while not block.has_class("tblock"):
            block = block.parent
        self.assertTrue(self.is_descendant(directory.by_id("t-models"), block))
        for path in self.site.rglob("*.html"):
            doc = Document(path.read_text(encoding="utf-8"))
            self.assertEqual(doc.errors, [], str(path))
            for node in doc.nodes:
                if node.tag != "a" or not node.attrs.get("href"):
                    continue
                href = urlsplit(node.attrs["href"])
                if href.scheme or href.netloc:
                    continue
                target = path.parent / unquote(href.path) if href.path else path
                self.assertTrue(target.is_file(), f"{path.name}: {node.attrs['href']}")
        self.assertFalse((self.site / "models" / "router.html").exists())

    def test_tab_dom_and_supported_filter_before_cap(self):
        text = (self.site / "index.html").read_text(encoding="utf-8")
        doc = Document(text)
        self.assertEqual(doc.errors, [])
        outer = doc.by_id("l-aa").parent
        self.assertTrue(outer.has_class("tab-group"))
        for key in ("l-aa", "l-b", "l-l", "l-v"):
            self.assertIs(doc.by_id(key).parent, outer)
        inner = doc.by_id("b-sup").parent
        self.assertTrue(inner.has_class("tab-group"))
        self.assertIs(inner.parent, doc.by_id("l-b"))
        self.assertIs(doc.by_id("b-all").parent, inner)
        for node in doc.nodes:
            if node.attrs.get("role") == "tab":
                panel = doc.by_id(node.attrs["aria-controls"])
                self.assertEqual(panel.attrs.get("role"), "tabpanel")
                self.assertEqual(panel.attrs.get("aria-labelledby"), node.attrs["id"])
                self.assertIs(panel.parent, node.parent.parent)
                self.assertEqual(node.attrs["aria-selected"] == "true", "hidden" not in panel.attrs)
        supported = doc.by_id("t-b")
        self.assertEqual(sum(n.tag == "tr" for n in next(n for n in supported.children if n.tag == "tbody").children), 100)
        self.assertIn("Model 110", supported.text)
        self.assertIn("Model 209", supported.text)
        self.assertNotIn("estimated", supported.text)
        self.assertIn("Showing 100 of 215 supported models", doc.by_id("b-sup").text)
        self.assertIn("Showing 100 of 325 models", doc.by_id("b-all").text)
        self.assertIn("querySelectorAll(':scope > [role=tab]')", text)

    def test_cross_view_model_links_and_callable_controls(self):
        allowed = {build_report.copy_id(m) for m in self.a["models"]} - {""}
        for path in list(self.site.rglob("*.html")) + [self.out / f"{self.a['stamp']}_report.html"]:
            doc = Document(path.read_text(encoding="utf-8"))
            controls = [n for n in doc.nodes if n.has_class("copy")]
            for node in controls:
                self.assertIn(node.text, allowed, str(path))
        doc = Document((self.site / "index.html").read_text(encoding="utf-8"))
        for key in ("t-b", "t-b2", "l-l", "l-v"):
            root = doc.by_id(key)
            self.assertTrue(any(n.tag == "a" and n.attrs.get("href", "").startswith("models/")
                                and self.is_descendant(n, root) for n in doc.nodes), key)
        for name in ("compare.html", "confidence.html"):
            doc = Document((self.site / name).read_text(encoding="utf-8"))
            self.assertTrue(any(n.tag == "a" and n.attrs.get("href", "").startswith("models/") for n in doc.nodes))

    @staticmethod
    def is_descendant(node, root):
        while node.parent:
            node = node.parent
            if node is root:
                return True
        return False

    def test_reliability_metadata_all_formats(self):
        stamp = self.a["stamp"]
        fields = ("run_id", "schema_version", "started_at", "source_health", "website_health", "churn")
        for path in (self.out / f"{stamp}_models.json", self.site / "data.json"):
            data = json.loads(path.read_text(encoding="utf-8"))
            for key in fields:
                self.assertEqual(data[key], self.a[key])
        for path in (self.out / f"{stamp}_report.html", self.site / "confidence.html"):
            doc = Document(path.read_text(encoding="utf-8"))
            values = [json.loads(n.text) for n in doc.nodes if n.tag == "pre"]
            self.assertTrue(all(n.parent.tag == "details" and "open" not in n.parent.attrs
                                for n in doc.nodes if n.tag == "pre"))
            for key in fields:
                self.assertIn(self.a[key], values)
        md = (self.out / f"{stamp}_summary.md").read_text(encoding="utf-8")
        for key in fields:
            value = re.search(r"## " + key + r"\n\n```json\n(.*?)\n```", md, re.S)[1]
            self.assertEqual(json.loads(value), self.a[key])
        from openpyxl import load_workbook
        wb = load_workbook(self.out / f"{stamp}_models.xlsx", read_only=True)
        try:
            self.assertEqual(len(wb.sheetnames), 15)
            summary = dict(wb["summary"].values)
            self.assertEqual(json.loads(summary["run_id"]), self.a["run_id"])
            self.assertEqual(json.loads(summary["schema_version"]), self.a["schema_version"])
            self.assertEqual(json.loads(summary["started_at"]), self.a["started_at"])
            health = dict(wb["Source_Health"].values)
            self.assertEqual(json.loads(health["aa.error"]), "<offline> & unavailable")
            self.assertEqual(json.loads(health["website_health.vals.components.pages.failed"]), 1)
            churn = dict(wb["Churn"].values)
            for field, value in self.a["churn"]["events"][0].items():
                path = f"events[0].{field}" + ("[0]" if field == "alternatives" else "")
                self.assertEqual(json.loads(churn[path]), value[0] if isinstance(value, list) else value)
            self.assertEqual(json.loads(churn["baselines.openrouter.run_id"]), "prior-run")
            self.assertEqual(json.loads(churn["daily.2026-09-26.disappeared"]), 1)
        finally:
            wb.close()

    def test_empty_sources_and_legacy_schema(self):
        with tempfile.TemporaryDirectory() as temp:
            a = fixture(count=0, metadata=False)
            out, site = generate(Path(temp), a)
            for path in site.glob("*.html"):
                self.assertEqual(Document(path.read_text(encoding="utf-8")).errors, [], path.name)
            index = (site / "index.html").read_text(encoding="utf-8")
            self.assertIn("BenchLM snapshot missing", index)
            self.assertEqual(list((site / "models").iterdir()), [])
            from openpyxl import load_workbook
            wb = load_workbook(out / f"{a['stamp']}_models.xlsx", read_only=True)
            try:
                self.assertEqual(len(wb.sheetnames), 13)
                self.assertNotIn("Churn", wb.sheetnames)
            finally:
                wb.close()
            data = json.loads((site / "data.json").read_text(encoding="utf-8"))
            self.assertNotIn("source_health", data)
            self.assertNotIn("churn", data)
            self.assertEqual((out / "old_models.json").read_text(), "keep")
            self.assertEqual((out / "old_site" / "keep.txt").read_text(), "keep")

    def test_model_coverage_with_empty_sources(self):
        with tempfile.TemporaryDirectory() as temp:
            a = fixture()
            a["views"] = {}
            a["source_health"] = {}
            a["website_health"] = {}
            a["churn"] = {"events": [], "baselines": {}, "daily": {}}
            out, site = generate(Path(temp), a)
            self.assertEqual(len(list((site / "models").glob("*.html"))), 325)
            for name in ("index.html", "models.html", "confidence.html"):
                self.assertEqual(Document((site / name).read_text(encoding="utf-8")).errors, [])
            data = json.loads((site / "data.json").read_text(encoding="utf-8"))
            self.assertEqual(data["churn"], a["churn"])
            for path in (site / "confidence.html", out / f"{a['stamp']}_report.html", out / f"{a['stamp']}_summary.md"):
                self.assertIn("No trusted baseline", path.read_text(encoding="utf-8"))
            from openpyxl import load_workbook
            wb = load_workbook(out / f"{a['stamp']}_models.xlsx", read_only=True)
            try:
                rows = dict(wb["Churn"].values)
                self.assertEqual(json.loads(rows["events"]), [])
                self.assertEqual(json.loads(rows["baselines"]), {})
                self.assertEqual(json.loads(rows["daily"]), {})
            finally:
                wb.close()

    def test_readable_health_churn_and_partial_summary(self):
        for path in (self.out / f"{self.a['stamp']}_report.html", self.site / "confidence.html"):
            doc = Document(path.read_text(encoding="utf-8"))
            section = doc.by_id("reliability")
            tables = [n for n in doc.nodes if n.tag == "table" and self.is_descendant(n, section)]
            self.assertEqual(len(tables), 3)
            for value in ("Source", "Status", "Count", "Fetched at", "Reason", "Baseline", "Coverage",
                          "2026-09-26T12:01:00Z", "prior-run", "<offline> & unavailable", "pages: attempted: 3; failed: 1"):
                self.assertIn(value, tables[0].text)
            for value in ("vendor/old", "catalog removal", "vendor/new", "partial"):
                self.assertIn(value, tables[1].text)
            self.assertIn("2026-09-26", tables[2].text)
            self.assertIn("disappeared: 1", tables[2].text)
        doc = Document((self.site / "index.html").read_text(encoding="utf-8"))
        summary = next(n for n in doc.nodes if n.has_class("source-summary"))
        self.assertIn("Partial coverage", summary.text)
        self.assertIn("aa: failed", summary.text)
        self.assertTrue(any(n.tag == "a" and n.attrs.get("href") == "confidence.html" for n in summary.children))

    def test_history_schema_nested_baselines_and_daily_rollups(self):
        a = {"source_health": {"openrouter": {"status": "complete", "count": 5}},
             "churn": {"trusted_route_history": True,
                       "source_health": {"openrouter": {"status": "complete", "complete": True, "compared": True,
                                                        "baseline": {"run_id": "prior", "ref": {"started_at": "yesterday"}}}},
                       "events": [{"type": "verified_free_paid", "provider": "openrouter", "id": "old/route",
                                   "canonical_id": "old-model", "alternatives": [{"provider": "zen", "id": "free/route"}],
                                   "unknown_coverage": True, "unknown_sources": ["nvidia"]}],
                       "daily": {"day": "2026-09-27", "sources": {"openrouter": {
                           "net_events": [], "observed_events": [{"type": "verified_free_paid"}],
                           "latest_complete": {"run_id": "current"}, "baseline": {"run_id": "prior"},
                           "coverage": [{"run_id": "current", "status": "complete"}]}}}}}
        doc = Document(build_report.reliability_html(a))
        self.assertEqual(doc.errors, [])
        tables = [n for n in doc.nodes if n.tag == "table"]
        self.assertIn("ref: started_at: yesterday", tables[0].text)
        self.assertIn("provider: zen; id: free/route", tables[1].text)
        self.assertIn("unknown_sources: nvidia", tables[1].text)
        self.assertIn("verified_free_paid: 1", tables[2].text)
        self.assertIn("No events", tables[2].text)
        self.assertNotIn("No trusted baseline", doc.root.text)

    def test_required_paths_and_identity_before_site_writes(self):
        for stage in (build_report.main, build_site.main):
            for kwargs in ({}, {"input_path": "missing"}, {"output_dir": "missing"}):
                with self.assertRaisesRegex(ValueError, "required"):
                    stage(**kwargs)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "input.json"
            a = fixture(count=0)
            source.write_text(json.dumps(a), encoding="utf-8")
            report = root / f"{a['stamp']}_models.json"
            for wrong in ({"stamp": "other", "run_id": a["run_id"]},
                          {"stamp": a["stamp"], "run_id": "other"}, {"stamp": a["stamp"]}):
                report.write_text(json.dumps(wrong), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "Mixed-run"):
                    build_site.main(source, root)
                self.assertFalse((root / f"{a['stamp']}_site").exists())
            for key in ("stamp", "run_id"):
                bad = {**a, key: "../bad"}
                with patch("builtins.open", return_value=io.StringIO(json.dumps(bad))) as opened:
                    with self.assertRaisesRegex(ValueError, "Invalid run ID"):
                        build_site.main(source, root)
                    self.assertEqual(opened.call_count, 1)

    def test_workbook_failure_propagates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "input.json"
            source.write_text(json.dumps(fixture(count=0)), encoding="utf-8")
            with patch("openpyxl.Workbook.save", side_effect=OSError("disk failure")):
                with self.assertRaisesRegex(OSError, "disk failure"):
                    build_report.main(source, root / "out")

    def test_explicit_stage_cli(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "input.json"
            a = fixture(count=0)
            source.write_text(json.dumps(a), encoding="utf-8")
            for module in (build_report, build_site):
                command = [sys.executable, "-B", module.__file__]
                missing = subprocess.run(command, cwd=root, capture_output=True, text=True)
                self.assertEqual(missing.returncode, 2, missing.stderr)
                result = subprocess.run(command + ["--input", str(source), "--output", str(root / "out")],
                                        cwd=root, capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((root / "out" / f"{a['stamp']}_site" / "index.html").is_file())


if __name__ == "__main__":
    unittest.main()
