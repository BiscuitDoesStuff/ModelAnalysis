"""Trace displayed scores and prices back to saved observations.

    python tools/audit_provenance.py <bundle> [--sample N]

Samples N displayed score/price cells per page and format (0 = all) from the
HTML dashboard, every site page, the XLSX *_provenance columns and the report
JSON, and follows each to its observation in the analysis JSON. Fails (exit 1)
on orphans (a shown value with no, or an unknown, obs_id), value mismatches and
wrong scales; reports missing versions/URLs and stale evidence as warnings.
"""
import argparse
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from analysis.scales import RANKING_UNIT  # noqa: E402

UNITS = {"score": RANKING_UNIT, "price": "usd-per-1m-blended-3to1"}
CELL = re.compile(r"<td class='(sc|cc)'(?: data-obs='([0-9a-f]+)')?[^>]*>(.*?)</td>", re.S)
NUMBER = re.compile(r"-?\d+(?:\.\d+)?(?:[eE]-?\d+)?")


def number(text):
    found = NUMBER.search(re.sub(r"<[^>]+>", " ", str(text)))
    return float(found.group()) if found else None


def sample(items, n):
    if not n or len(items) <= n:
        return items
    step = len(items) / n
    return [items[int(i * step)] for i in range(n)]


def cells(bundle, run_id):
    """(format, page, field, obs_id or None, displayed value or None)."""
    reports = bundle / "reports"
    pages = [reports / f"{run_id}_report.html"] + sorted((reports / f"{run_id}_site").rglob("*.html"))
    for page in pages:
        if not page.exists():
            continue
        fmt = "dashboard" if page.name.endswith("_report.html") else "site"
        found = []
        for cls, oid, text in CELL.findall(page.read_text(encoding="utf-8")):
            found.append((fmt, page.relative_to(bundle).as_posix(), "score" if cls == "sc" else "price",
                          oid or None, number(text)))
        yield page.relative_to(bundle).as_posix(), found
    xlsx = reports / f"{run_id}_models.xlsx"
    if xlsx.exists():
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = list(next(rows, []) or [])
            pairs = [(f, header.index(v), header.index(p)) for f, v, p in
                     (("score", "score", "score_provenance"), ("price", "cost_per_1M", "cost_provenance"))
                     if v in header and p in header]
            if not pairs:
                continue
            found = []
            for row in rows:
                for field, vi, pi in pairs:
                    prov = row[pi] or ""
                    oid = prov.rsplit(" · ", 1)[-1] if prov else None
                    found.append(("xlsx", f"xlsx:{ws.title}", field, oid, row[vi] if isinstance(row[vi], (int, float)) else None))
            sheets.append((f"xlsx:{ws.title}", found))
        wb.close()
        yield from sheets
    report = json.loads((reports / f"{run_id}_models.json").read_text(encoding="utf-8"))
    found = []
    for m in (report.get("models_by_slug") or {}).values():
        for field, key in (("score", "score"), ("price", "cost_blended")):
            oid = ((m.get("provenance") or {}).get(field) or {}).get("obs_id")
            found.append(("json", "report json", field, oid, m.get(key)))
    yield "report json", found


def audit(bundle, n=25):
    bundle = Path(bundle)
    run_id = json.loads((bundle / "manifest.json").read_text(encoding="utf-8"))["run_id"]
    a = json.loads((bundle / "analysis" / f"{run_id}_analysis.json").read_text(encoding="utf-8"))
    observations = {o["obs_id"]: o for o in a.get("observations", [])}
    day = a.get("day", "")
    result = {"checked": 0, "orphans": [], "mismatches": [], "wrong_scale": [],
              "missing_version": [], "missing_url": [], "stale": [], "per_format": {}}
    for page, found in cells(bundle, run_id):
        shown = [c for c in found if c[3] or c[4] is not None]
        for fmt, where, field, oid, value in sample(shown, n):
            result["checked"] += 1
            result["per_format"][fmt] = result["per_format"].get(fmt, 0) + 1
            label = f"{where} {field} {oid or '(no obs_id)'}"
            o = observations.get(oid) if oid else None
            if o is None:
                result["orphans"].append(label)
                continue
            if value is None or not isinstance(o.get("value"), (int, float)) or \
                    abs(float(value) - float(o["value"])) > 1e-9 * max(1.0, abs(float(o["value"]))):
                result["mismatches"].append(f"{label}: shown {value} vs observed {o.get('value')}")
            if o.get("unit") != UNITS[field]:
                result["wrong_scale"].append(f"{label}: unit {o.get('unit')}")
            if field == "score" and not o.get("version"):
                result["missing_version"].append(f"{label}: {o.get('version_unavailable_reason', 'no reason given')}")
            if not o.get("url") and not o.get("url_unavailable_reason"):
                result["missing_url"].append(label)
            if o.get("expires_at") and day and o["expires_at"] < day:
                result["stale"].append(f"{label}: expired {o['expires_at']}")
    return result


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("bundle")
    p.add_argument("--sample", type=int, default=25, help="cells per page/format; 0 = all")
    args = p.parse_args()
    r = audit(args.bundle, args.sample)
    print(f"checked {r['checked']} displayed cells: " + ", ".join(f"{k} {v}" for k, v in sorted(r["per_format"].items())))
    for key, severity in (("orphans", "FAIL"), ("mismatches", "FAIL"), ("wrong_scale", "FAIL"),
                          ("missing_version", "warn"), ("missing_url", "warn"), ("stale", "warn")):
        print(f"{severity if r[key] else 'ok  '} {key}: {len(r[key])}")
        for line in r[key][:10]:
            print("     " + line)
    sys.exit(1 if r["orphans"] or r["mismatches"] or r["wrong_scale"] else 0)


if __name__ == "__main__":
    main()
