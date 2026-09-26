"""3. Report — 9-section OCF views: MD + XLSX + JSON + HTML per run. On-use. Keeps latest only."""
import json, os, glob, html as _html

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REP = os.path.join(ROOT, "reports")
os.makedirs(REP, exist_ok=True)

MD_CAP = 20
TIERS = ["max", "high", "medium"]
TIER_LABEL = {"max": "Max (50+)", "high": "High (40+)", "medium": "Medium/General (30+)"}
VARIANTS = {"OCF": set(), "OF": {"C"}, "CF": {"O"}, "F": {"O", "C"}}
OCF = {"O", "C", "F"}


def opencode_ids(or_id):
    if not or_id:
        return ""
    return or_id if or_id.startswith("openrouter/") else f"openrouter/{or_id}"


def hier_key(r):
    s = r.get("score")
    x = r.get("ratio")
    return (-(s if s is not None else -1), -(x if x is not None else -1))


def free_status_of(m):
    fs = m.get("free_status")
    if fs in ("verified", "provisional-l1", "provisional-l0", "none"):
        return fs
    return "verified" if m.get("free") else "none"


def is_provisional(m):
    return free_status_of(m) in ("provisional-l1", "provisional-l0")


def has_callable(m):
    if m.get("or_id"):
        return True
    provs = set(m.get("providers", []))
    return bool(provs & {"openai", "anthropic", "openrouter", "groq", "cerebras",
                         "nvidia", "zenmux", "zen"})


def groups_display(m):
    g = "".join(m.get("groups", [])) or ""
    if is_provisional(m):
        g = (g + "F?") if g else "F?"
    return g or "–"


def pct(vals, p):
    if not vals:
        return None
    s = sorted(vals)
    k = (len(s) - 1) * p
    f, c = int(k), min(int(k) + 1, len(s) - 1)
    return s[f] + (s[c] - s[f]) * (k - f)


def prune_reports(keep_stamps):
    pats = ["*_summary.md", "*_models.json", "*_models.xlsx", "*_report.html", "*_churn_alert.md"]
    files = []
    for p in pats:
        files += glob.glob(os.path.join(REP, p))
    for f in files:
        bn = os.path.basename(f)
        if not any(bn.startswith(s) for s in keep_stamps):
            os.remove(f)
            print(f"pruned report {os.path.basename(f)}")


def build_sections(a):
    models = a.get("models", [])
    routers = [m["id"] for m in models if m.get("router")]
    models = [m for m in models if not m.get("router")]
    ocf = [m for m in models
           if set(m.get("groups", [])) & OCF or is_provisional(m)]

    all_intel = sorted(models, key=lambda r: (-(r["score"] if r["score"] is not None else -1)))
    costed = sorted([m for m in models if m.get("cost_blended") is not None],
                    key=lambda r: r["cost_blended"])
    uncosted = [m for m in models if m.get("cost_blended") is None]
    paid_ratio = sorted([m for m in models
                         if free_status_of(m) == "none" and m.get("ratio") is not None],
                        key=lambda r: -r["ratio"])
    free_block = sorted([m for m in models
                         if free_status_of(m) == "verified" and m.get("score") is not None],
                        key=lambda r: -r["score"])
    provisional_block = sorted([m for m in models
                                if is_provisional(m) and m.get("score") is not None],
                               key=lambda r: -r["score"])
    free_unscored = [m for m in models
                     if free_status_of(m) == "verified" and m.get("score") is None]
    provisional_unscored = [m for m in models
                            if is_provisional(m) and m.get("score") is None]
    unratable = [m for m in models
                 if free_status_of(m) == "none" and m.get("ratio") is None]

    ocf_intel = [m for m in all_intel if m in ocf]
    ocf_costed = [m for m in costed if m in ocf]
    ocf_uncosted = [m for m in uncosted if m in ocf]
    ocf_ratio = [m for m in paid_ratio if m in ocf]
    ocf_free = [m for m in free_block if m in ocf]
    ocf_provisional = [m for m in provisional_block if m in ocf]
    ocf_free_unscored = [m for m in free_unscored if m in ocf]
    ocf_provisional_unscored = [m for m in provisional_unscored if m in ocf]
    ocf_unratable = [m for m in unratable if m in ocf]

    stack = {}
    for t in TIERS:
        rows = sorted([m for m in ocf if m.get("tier") == t and has_callable(m)],
                      key=hier_key)
        present = set().union(*[set(m.get("groups", [])) for m in rows]) if rows else set()
        stack[t] = {"rows": rows, "gaps": sorted(OCF - present)}

    practical = []
    for t in TIERS:
        for v, off in VARIANTS.items():
            rows = sorted([m for m in stack[t]["rows"] if not (set(m.get("groups", [])) & off)],
                          key=hier_key)
            practical.append({"tier": t, "variant": v,
                              "winner": rows[0] if rows else None,
                              "runner_up": rows[1] if len(rows) > 1 else None})

    elig = [m for m in ocf if free_status_of(m) == "none" and m.get("score") is not None
            and m.get("cost_blended") is not None]
    q = {"score_q1": pct([m["score"] for m in elig], 0.25),
         "score_q3": pct([m["score"] for m in elig], 0.75),
         "cost_q1": pct([m["cost_blended"] for m in elig], 0.25),
         "cost_q3": pct([m["cost_blended"] for m in elig], 0.75)}
    outliers = {"bargains": [], "overpriced": [], "free_gems": []}
    if elig and q["score_q3"] is not None:
        outliers["bargains"] = [m for m in elig
                                if m["score"] >= q["score_q3"] and m["cost_blended"] <= q["cost_q1"]]
        outliers["overpriced"] = [m for m in elig
                                  if m["score"] <= q["score_q1"] and m["cost_blended"] >= q["cost_q3"]]
    outliers["free_gems"] = [m for m in ocf
                             if (free_status_of(m) in ("verified", "provisional-l1", "provisional-l0"))
                             and (m.get("score") or 0) >= 40 and has_callable(m)]

    return {"all_intel": all_intel, "costed": costed, "uncosted": uncosted,
            "paid_ratio": paid_ratio, "free_block": free_block,
            "provisional_block": provisional_block,
            "free_unscored": free_unscored,
            "provisional_unscored": provisional_unscored, "unratable": unratable,
            "ocf_intel": ocf_intel, "ocf_costed": ocf_costed, "ocf_uncosted": ocf_uncosted,
            "ocf_ratio": ocf_ratio, "ocf_free": ocf_free,
            "ocf_provisional": ocf_provisional,
            "ocf_free_unscored": ocf_free_unscored,
            "ocf_provisional_unscored": ocf_provisional_unscored,
            "ocf_unratable": ocf_unratable,
            "stack": stack, "practical": practical, "outliers": outliers,
            "quartiles": q, "ocf_count": len(ocf), "model_count": len(models),
            "routers": routers}


def row_md(m):
    g = groups_display(m) if groups_display(m) != "–" else "-"
    s = m["score"] if m.get("score") is not None else "unscored"
    c = m["cost_blended"] if m.get("cost_blended") is not None else "cost-unknown"
    r = m["ratio"] if m.get("ratio") is not None else "-"
    oc = f" → `{opencode_ids(m['or_id'])}`" if m.get("or_id") else ""
    return f"- `{m['id']}` [{g}] — score {s} — ${c}/1M — ratio {r}{oc}\n"


def esc(v):
    return _html.escape("" if v is None else str(v))


def fmt_cost(m):
    return "unknown" if m.get("cost_blended") is None else f"${m['cost_blended']}/1M"


def fmt_score(m):
    return "unscored" if m.get("score") is None else str(m["score"])


def html_table(rows, note=""):
    h = ['<input class="search" placeholder="Filter…" oninput="filterRows(this)">']
    if note:
        h.append(f'<p class="note">{esc(note)}</p>')
    h.append('<div class="twrap"><table><thead><tr><th>Model</th><th>Groups</th>'
             '<th>Score</th><th>Cost</th><th>Ratio</th><th>OpenCode ID <span class="hint">(click to copy)</span></th>'
             '<th>Sources</th></tr></thead><tbody>')
    for m in rows:
        oc = opencode_ids(m.get("or_id", ""))
        h.append("<tr><td><code>" + esc(m["id"]) + "</code>" +
                 (f"<br><span class='nm'>{esc(m.get('name', ''))}</span>" if m.get("name") else "") +
                 "</td><td>" + esc(groups_display(m)) + "</td><td>" + esc(fmt_score(m)) +
                 "</td><td>" + esc(fmt_cost(m)) + "</td><td>" +
                 esc(m["ratio"] if m.get("ratio") is not None else "–") + "</td><td>" +
                 (f"<code class='copy' onclick=\"copyId(this)\" title='click to copy'>{esc(oc)}</code>" if oc else "–") +
                 "</td><td>" + esc(",".join(m.get("providers", []))) + "</td></tr>")
    h.append("</tbody></table></div>")
    return "".join(h)


def build_html(a, s, stamp):
    th = a.get("thresholds", {"max": 50, "high": 40, "medium": 30})
    fsc = a.get("free_status_counts", {})
    chips = (f"<span class='chip'>Models {s['model_count']}</span>"
             f"<span class='chip'>OCF {s['ocf_count']}</span>"
             f"<span class='chip'>Scored {sum(1 for m in a.get('models', []) if m.get('score') is not None)}</span>"
             f"<span class='chip'>Max {len(s['stack']['max']['rows'])}</span>"
             f"<span class='chip'>High {len(s['stack']['high']['rows'])}</span>"
             f"<span class='chip'>Medium {len(s['stack']['medium']['rows'])}</span>"
             f"<span class='chip'>F verified {fsc.get('verified', sum(1 for m in a.get('models', []) if free_status_of(m) == 'verified'))}</span>"
             f"<span class='chip'>F? prov {fsc.get('provisional-l1', 0) + fsc.get('provisional-l0', 0)}</span>"
             f"<span class='chip dim'>Routers {len(s['routers'])} excluded</span>")

    def sec(tid, title, body):
        return f"<section id='t-{tid}' class='tab'><h2>{esc(title)}</h2>{body}</section>"

    tabs = [("overview", "Overview"), ("all-intel", "All · Intelligence"),
            ("all-cost", "All · Cost"), ("all-ratio", "All · Ratio"),
            ("ocf-intel", "OCF · Intelligence"), ("ocf-cost", "OCF · Cost"),
            ("ocf-ratio", "OCF · Ratio"), ("stack", "Stack"),
            ("practical", "Practical picks"), ("outliers", "Outliers")]
    nav = "".join(f"<button data-t='t-{tid}' onclick='showTab(this)'>{t}</button>" for tid, t in tabs)

    ov = f"<p>Snapshot <b>{esc(stamp)}</b> · thresholds {th['max']}/{th['high']}/{th['medium']} " \
         "(below 30 in All views only) · free models never enter ratios, ranked by score instead. " \
         "[F?] = provisional free (AA $0, billing unverified; exact level in JSON/XLSX).</p>"
    ch = a.get("free_churn")
    if ch and ch.get("prev_day"):
        ov += (f"<p class='note'>Free-status churn vs {esc(ch['prev_day'])}: "
               f"→paid {ch.get('flipped_to_paid_total', 0)} · →free {ch.get('flipped_to_free_total', 0)} · "
               f"level-changed {ch.get('level_changed_total', 0)} · disappeared {ch.get('disappeared_total', 0)} · "
               f"new {ch.get('new_total', 0)}.</p>")
    ov += "<div class='cards'>"
    for t in TIERS:
        rows = s["stack"][t]["rows"]
        top = rows[0] if rows else None
        gaps = f"<span class='gap'>gaps: {','.join(s['stack'][t]['gaps'])}</span>" if s["stack"][t]["gaps"] else ""
        ov += (f"<div class='card'><h3>{TIER_LABEL[t]}</h3>"
               f"<p class='bign'>{len(rows)} models {gaps}</p>" +
               (f"<p>Top: <code>{esc(top['id'])}</code> ({fmt_score(top)}, {fmt_cost(top)})</p>" if top else "<p>—</p>") +
               "</div>")
    ov += "</div><h3>Practical winners</h3><table><thead><tr><th>Tier</th><th>OCF</th><th>OF</th><th>CF</th><th>F-only</th></tr></thead><tbody>"
    for t in TIERS:
        cells = []
        for v in ("OCF", "OF", "CF", "F"):
            p = next(x for x in s["practical"] if x["tier"] == t and x["variant"] == v)
            cells.append(f"<code>{esc(p['winner']['id'])}</code>" if p["winner"] else "<span class='gap'>gap</span>")
        ov += f"<tr><td>{TIER_LABEL[t]}</td><td>{cells[0]}</td><td>{cells[1]}</td><td>{cells[2]}</td><td>{cells[3]}</td></tr>"
    ov += "</tbody></table>"

    stack_html = ""
    for t in TIERS:
        gaps = f" <span class='gap'>gaps: {','.join(s['stack'][t]['gaps'])}</span>" if s["stack"][t]["gaps"] else ""
        stack_html += f"<h3>{TIER_LABEL[t]}{gaps}</h3>" + html_table(s["stack"][t]["rows"])
    prac = ("<table><thead><tr><th>Tier</th><th>Variant</th><th>Winner</th><th>Runner-up</th></tr></thead><tbody>" +
            "".join(f"<tr><td>{TIER_LABEL[p['tier']]}</td><td>{p['variant']}</td>"
                     f"<td>{('<code>' + esc(p['winner']['id']) + '</code> ' + esc(fmt_score(p['winner'])) + ' / ' + esc(p['winner']['ratio'] if p['winner']['ratio'] is not None else '–')) if p['winner'] else '<span class=gap>gap</span>'}</td>"
                     f"<td>{('<code>' + esc(p['runner_up']['id']) + '</code>') if p['runner_up'] else '–'}</td></tr>"
                     for p in s["practical"]) + "</tbody>")
    out = "".join(f"<h3>{t}</h3>" + html_table(s["outliers"][k]) for k, t in
                  [("bargains", "Bargains — top-quartile score, bottom-quartile cost"),
                   ("overpriced", "Overpriced — bottom-quartile score, top-quartile cost"),
                   ("free_gems", "Free gems — verified/provisional free, score 40+")])

    body = (sec("overview", "Overview", ov) +
            sec("all-intel", "1 · All Data — Intelligence", html_table(s["all_intel"])) +
            sec("all-cost", "2 · All Data — Cost", html_table(s["costed"]) +
                (f"<p class='note'>Cost unknown: {len(s['uncosted'])} — " + esc(", ".join(m["id"] for m in s["uncosted"][:50])) + "</p>" if s["uncosted"] else "")) +
            sec("all-ratio", "3 · All Data — Ratio (paid) + free by score",
                "<h3>Paid by ratio</h3>" + html_table(s["paid_ratio"]) +
                "<h3>Verified free by score</h3>" + html_table(s["free_block"]) +
                "<h3>Provisional free [F?] by score</h3>" + html_table(s["provisional_block"])) +
            sec("ocf-intel", "4 · OCF — Intelligence", html_table(s["ocf_intel"])) +
            sec("ocf-cost", "5 · OCF — Cost", html_table(s["ocf_costed"])) +
            sec("ocf-ratio", "6 · OCF — Ratio + free by score",
                "<h3>Paid by ratio</h3>" + html_table(s["ocf_ratio"]) +
                "<h3>Verified free by score</h3>" + html_table(s["ocf_free"]) +
                "<h3>Provisional free [F?] by score</h3>" + html_table(s["ocf_provisional"])) +
            sec("stack", "7 · OCF — Optimized stack", stack_html) +
            sec("practical", "8 · OCF — Practical picks", prac) +
            sec("outliers", "9 · OCF — Outliers", out))

    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>ModelAnalysis — {esc(stamp)}</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;background:#020617;color:#e2e8f0;margin:0;padding:24px;max-width:1200px}}
h1{{font-size:24px}}h2{{color:#7dd3fc}}h3{{color:#bae6fd}}
.chip{{display:inline-block;background:#082f49;border:1px solid #38bdf8;border-radius:12px;padding:3px 12px;margin:2px;font-size:13px}}
.dim{{opacity:.6}}nav{{position:sticky;top:0;background:#020617;padding:10px 0;z-index:5}}
nav button{{background:#1e293b;color:#e2e8f0;border:1px solid #38bdf8;border-radius:8px;padding:6px 12px;margin:2px;cursor:pointer}}
nav button.on{{background:#0369a1}}.tab{{display:none}}.tab.on{{display:block}}
table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{border:1px solid #334155;padding:6px 8px;text-align:left;vertical-align:top}}
th{{background:#0f172a}}tr:nth-child(even){{background:#0b1220}}code{{color:#7dd3fc}}
.copy{{cursor:pointer;border-bottom:1px dotted #38bdf8}}.nm{{color:#94a3b8;font-size:12px}}
.gap{{color:#fbbf24;font-weight:700}}.note{{color:#94a3b8}}.hint{{font-weight:400;font-size:11px;color:#94a3b8}}
.search{{width:280px;padding:6px 10px;margin:8px 0;background:#0f172a;color:#e2e8f0;border:1px solid #38bdf8;border-radius:8px}}
.twrap{{overflow-x:auto}}.cards{{display:flex;gap:12px;flex-wrap:wrap}}.card{{background:#0f172a;border:1px solid #38bdf8;border-radius:10px;padding:12px 16px;min-width:200px}}
.bign{{font-size:15px}}
</style></head><body>
<h1>ModelAnalysis — {esc(stamp)}</h1>
<div>{chips}</div><nav>{nav}</nav>{body}
<script>
function showTab(b){{document.querySelectorAll('nav button').forEach(x=>x.classList.remove('on'));
document.querySelectorAll('.tab').forEach(x=>x.classList.remove('on'));
b.classList.add('on');document.getElementById(b.dataset.t).classList.add('on');}}
function filterRows(inp){{const q=inp.value.toLowerCase();
inp.parentElement.querySelectorAll('tbody tr').forEach(tr=>{{
tr.style.display=tr.textContent.toLowerCase().includes(q)?'':'none';}});}}
function copyId(el){{navigator.clipboard.writeText(el.textContent).then(()=>{{
el.style.color='#4ade80';setTimeout(()=>el.style.color='',800);}});}}
document.querySelector('nav button').classList.add('on');
document.querySelector('.tab').classList.add('on');
</script></body></html>"""


def main():
    afiles = sorted(glob.glob(os.path.join(ROOT, "analysis", "*_analysis.json")), key=os.path.getmtime)
    if not afiles:
        print("run analysis first")
        return
    with open(afiles[-1], encoding="utf-8") as f:
        a = json.load(f)
    stamp = a.get("stamp", a.get("day", "unknown"))
    th = a.get("thresholds", {"max": 50, "high": 40, "medium": 30})
    s = build_sections(a)

    L = [f"# ModelAnalysis — {stamp} (on-use)\n",
         f"Models {s['model_count']} | OCF {s['ocf_count']} | "
         f"Scored {sum(1 for m in a.get('models', []) if m.get('score') is not None)} | "
         f"Costed {sum(1 for m in a.get('models', []) if m.get('cost_blended') is not None)} | "
         f"Tiers Max {len(s['stack']['max']['rows'])} / High {len(s['stack']['high']['rows'])} / "
         f"Medium {len(s['stack']['medium']['rows'])} | "
         f"Thresholds {th['max']}/{th['high']}/{th['medium']} | <30 excluded from stack only | "
         f"F verified {sum(1 for m in a.get('models', []) if free_status_of(m) == 'verified')} | "
         f"F? provisional L1/L0 "
         f"{sum(1 for m in a.get('models', []) if free_status_of(m) == 'provisional-l1')}/"
         f"{sum(1 for m in a.get('models', []) if free_status_of(m) == 'provisional-l0')} | "
         f"Routers {len(s['routers'])} excluded\n",
         "\n## 1. All Data — Intelligence\n"]
    L += [row_md(m) for m in s["all_intel"][:MD_CAP]]
    if len(s["all_intel"]) > MD_CAP:
        L.append(f"_…and {len(s['all_intel']) - MD_CAP} more (see XLSX/JSON)_\n")
    L += ["\n## 2. All Data — Cost\n"]
    L += [row_md(m) for m in s["costed"][:MD_CAP]]
    if s["uncosted"]:
        L.append(f"\n_cost-unknown ({len(s['uncosted'])}), e.g: " +
                 ", ".join(f"`{m['id']}`" for m in s["uncosted"][:MD_CAP]) + "_\n")
    L += ["\n## 3. All Data — Intelligence/Cost Ratio (paid; verified-free by score; provisional [F?] by score)\n"]
    L += [row_md(m) for m in s["paid_ratio"][:MD_CAP]]
    L += ["\n_Verified free by score_\n"]
    L += [row_md(m) for m in s["free_block"][:MD_CAP]]
    L += ["\n_Provisional free [F?] by score (AA $0, billing unverified)_\n"]
    L += [row_md(m) for m in s["provisional_block"][:MD_CAP]]
    if s["free_unscored"] or s["provisional_unscored"]:
        L.append(f"_Free unscored verified ({len(s['free_unscored'])})"
                 + (", ".join(f"`{m['id']}`" for m in s["free_unscored"][:MD_CAP]) if s["free_unscored"] else "—") +
                 f" | provisional unscored ({len(s['provisional_unscored'])})_\n")
    L += ["\n## 4. OCF — Intelligence\n"]
    L += [row_md(m) for m in s["ocf_intel"][:MD_CAP]]
    L += ["\n## 5. OCF — Cost\n"]
    L += [row_md(m) for m in s["ocf_costed"][:MD_CAP]]
    if s["ocf_uncosted"]:
        L.append(f"\n_cost-unknown ({len(s['ocf_uncosted'])})_\n")
    L += ["\n## 6. OCF — Intelligence/Cost Ratio\n"]
    L += [row_md(m) for m in s["ocf_ratio"][:MD_CAP]]
    L += ["\n_Verified free by score_\n"]
    L += [row_md(m) for m in s["ocf_free"][:MD_CAP]]
    L += ["\n_Provisional free [F?] by score_\n"]
    L += [row_md(m) for m in s["ocf_provisional"][:MD_CAP]]
    L += ["\n## 7. OCF — Optimized stack\n"]
    for t in TIERS:
        gaps = f" — gaps: {','.join(s['stack'][t]['gaps'])}" if s["stack"][t]["gaps"] else ""
        L.append(f"\n### {TIER_LABEL[t]}{gaps}\n")
        L += [row_md(m) for m in s["stack"][t]["rows"][:MD_CAP]]
    L += ["\n## 8. OCF — Practical picks (winner + runner-up per tier × variant)\n",
          "| Tier | Variant | Winner | Runner-up |\n|---|---|---|---|\n"]
    for p in s["practical"]:
        w = f"`{p['winner']['id']}` ({p['winner']['score']}/{p['winner']['ratio']})" if p["winner"] else "— (gap)"
        u = f"`{p['runner_up']['id']}`" if p["runner_up"] else "—"
        L.append(f"| {TIER_LABEL[p['tier']]} | {p['variant']} | {w} | {u} |\n")
    L += ["\n## 9. OCF — Outliers\n"]
    for cls, title in [("bargains", "Bargains (top-quartile score, bottom-quartile cost)"),
                       ("overpriced", "Overpriced (bottom-quartile score, top-quartile cost)"),
                       ("free_gems", "Free gems (verified/provisional free, score 40+)")]:
        L.append(f"\n### {title}\n")
        L += [row_md(m) for m in s["outliers"][cls][:MD_CAP]] or ["_(none)_\n"]
    L.append("\n_Free excluded from ratio (infinite); ratios use " +
             f"{a.get('cost_method', 'blended $/1M')}. OCF = OpenAI + Claude + strict-free + provisional [F?] " +
             "(AA $0, billing unverified; stack/practical need a callable or_id/native ID, " +
             "AA-only rows stay in Intel/Cost/Ratio; tier filled only by provisional still flags gap F(verified)). " +
             "Per-model OpenCode compat not verified._\n")
    ch = a.get("free_churn")
    if ch and ch.get("prev_day"):
        L.append(f"\n_Churn vs {ch['prev_day']}: →paid {ch.get('flipped_to_paid_total', 0)} " +
                 (", ".join(f"`{i}`" for i in ch.get('flipped_to_paid', [])[:MD_CAP]) if ch.get('flipped_to_paid') else "—") +
                 f" · →free {ch.get('flipped_to_free_total', 0)}" +
                 (", ".join(f"`{i}`" for i in ch.get('flipped_to_free', [])[:MD_CAP]) if ch.get('flipped_to_free') else "") +
                 f" · disappeared {ch.get('disappeared_total', 0)} · new {ch.get('new_total', 0)}._\n")
    with open(os.path.join(REP, f"{stamp}_summary.md"), "w", encoding="utf-8") as f:
        f.write("".join(L))

    slim = lambda m: [m["id"], groups_display(m), m["score"], m["cost_blended"],
                      m["ratio"], opencode_ids(m["or_id"]), ",".join(m["providers"]),
                      free_status_of(m)]
    scored = sorted(m["score"] for m in a.get("models", []) if m.get("score") is not None)
    dist = {"scored": len(scored)}
    if scored:
        dist.update({"p10": round(pct(scored, 0.10), 2), "p50": round(pct(scored, 0.50), 2),
                     "p90": round(pct(scored, 0.90), 2), "max": max(scored)})
    full = {"stamp": stamp, "day": a.get("day", stamp[:10]), "thresholds": th,
            "cost_method": a.get("cost_method", ""), "quartiles": s["quartiles"],
            "score_dist": dist, "routers_excluded": s["routers"],
            "collisions": a.get("collisions", []),
            "free_status_counts": a.get("free_status_counts", {}),
            "free_churn": a.get("free_churn", {}),
            "all_intel": s["all_intel"], "all_cost": s["costed"], "all_cost_unknown": s["uncosted"],
            "all_ratio_paid": s["paid_ratio"], "all_ratio_free_by_score": s["free_block"],
            "all_ratio_verified_free_by_score": s["free_block"],
            "all_ratio_provisional_free_by_score": s["provisional_block"],
            "all_ratio_unratable": s["unratable"] + s["free_unscored"] + s["provisional_unscored"],
            "ocf_intel": s["ocf_intel"], "ocf_cost": s["ocf_costed"],
            "ocf_cost_unknown": s["ocf_uncosted"], "ocf_ratio_paid": s["ocf_ratio"],
            "ocf_ratio_free_by_score": s["ocf_free"],
            "ocf_ratio_verified_free_by_score": s["ocf_free"],
            "ocf_ratio_provisional_free_by_score": s["ocf_provisional"],
            "ocf_stack": {t: {"rows": s["stack"][t]["rows"], "gaps": s["stack"][t]["gaps"]} for t in TIERS},
            "ocf_practical": s["practical"], "ocf_outliers": s["outliers"]}
    with open(os.path.join(REP, f"{stamp}_models.json"), "w", encoding="utf-8") as f:
        json.dump(full, f, indent=1)

    try:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "summary"
        ws.append(["stamp", stamp])
        ws.append(["models", s["model_count"]])
        ws.append(["ocf", s["ocf_count"]])
        fsc = a.get("free_status_counts", {})
        ws.append(["free_verified", fsc.get("verified", "")])
        ws.append(["provisional_l1", fsc.get("provisional-l1", "")])
        ws.append(["provisional_l0", fsc.get("provisional-l0", "")])
        ws.append(["total_groq", a.get("total_groq", "")])
        ws.append(["total_cerebras", a.get("total_cerebras", "")])
        ch = a.get("free_churn", {})
        if ch and ch.get("prev_day"):
            ws.append(["churn_vs", ch.get("prev_day", "")])
            ws.append(["churn_to_paid", ch.get("flipped_to_paid_total", 0)])
            ws.append(["churn_to_free", ch.get("flipped_to_free_total", 0)])
            ws.append(["churn_disappeared", ch.get("disappeared_total", 0)])
            ws.append(["churn_new", ch.get("new_total", 0)])
        for t in TIERS:
            ws.append([f"stack_{t}", len(s["stack"][t]["rows"])])
            ws.append([f"gaps_{t}", ",".join(s["stack"][t]["gaps"])])
        H = ["id", "groups", "score", "cost_per_1M", "ratio", "opencode_id", "providers", "free_status"]
        tabs = {"All_Intel": s["all_intel"], "All_Cost": s["costed"] + s["uncosted"],
                "All_Ratio": s["paid_ratio"] + s["free_block"] + s["provisional_block"] + s["unratable"] + s["free_unscored"] + s["provisional_unscored"],
                "OCF_Intel": s["ocf_intel"], "OCF_Cost": s["ocf_costed"] + s["ocf_uncosted"],
                "OCF_Ratio": s["ocf_ratio"] + s["ocf_free"] + s["ocf_provisional"] + s["ocf_unratable"] + s["ocf_free_unscored"] + s["ocf_provisional_unscored"],
                "OCF_Outliers": s["outliers"]["bargains"] + s["outliers"]["overpriced"] + s["outliers"]["free_gems"]}
        for name, rows in tabs.items():
            w = wb.create_sheet(name); w.append(H)
            for m in rows:
                w.append(slim(m))
        w = wb.create_sheet("OCF_Stack"); w.append(["tier"] + H + ["gap"])
        for t in TIERS:
            for m in s["stack"][t]["rows"]:
                w.append([t] + slim(m) + [""])
            for g in s["stack"][t]["gaps"]:
                w.append([t, f"GAP:{g}", "", "", "", "", "", "", "", "gap"])
        w = wb.create_sheet("OCF_Practical")
        w.append(["tier", "variant", "winner", "winner_score", "winner_ratio", "runner_up"])
        for p in s["practical"]:
            w.append([p["tier"], p["variant"],
                      p["winner"]["id"] if p["winner"] else "GAP",
                      p["winner"]["score"] if p["winner"] else None,
                      p["winner"]["ratio"] if p["winner"] else None,
                      p["runner_up"]["id"] if p["runner_up"] else None])
        wb.save(os.path.join(REP, f"{stamp}_models.xlsx"))
        x = "xlsx ok"
    except Exception as e:
        x = f"xlsx skipped: {e}"
    with open(os.path.join(REP, f"{stamp}_report.html"), "w", encoding="utf-8") as f:
        f.write(build_html(a, s, stamp))
    prune_reports([os.path.basename(f).replace("_analysis.json", "") for f in afiles[-1:]])
    print(f"report {stamp} written ({x} + html)")


if __name__ == "__main__":
    main()
